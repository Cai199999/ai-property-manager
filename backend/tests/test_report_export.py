import os
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.core.exceptions import ReportNameInvalidError, ReportNotFoundError
from app.services.chart_meta import infer_chart_meta
from app.services.report_export import ExcelReportExporter


def _settings(tmp_path: Path, **overrides: object) -> Settings:
    values: dict[str, object] = {
        "report_export_enabled": True,
        "report_export_dir": tmp_path / "exports",
        "report_export_ttl_seconds": 3600,
    }
    values.update(overrides)
    return Settings(_env_file=None, **values)


def test_infer_chart_meta_prefers_line_for_date_axis() -> None:
    columns = [
        {"name": "日期", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
    ]
    rows = [
        {"日期": "2026-09-01", "工单量": 3},
        {"日期": "2026-09-02", "工单量": 5},
    ]

    meta = infer_chart_meta(columns, rows)

    assert meta.type == "line"
    assert meta.x == "日期"
    assert meta.y == ["工单量"]
    assert meta.series[0].unit == "单"


def test_infer_chart_meta_uses_pie_for_small_category_count() -> None:
    columns = [
        {"name": "所属区域", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
    ]
    rows = [
        {"所属区域": "朝阳区", "工单量": 10},
        {"所属区域": "海淀区", "工单量": 4},
    ]

    meta = infer_chart_meta(columns, rows)

    assert meta.type == "pie"
    assert meta.title == "工单量按所属区域"


def test_infer_chart_meta_honours_explicit_chart_request() -> None:
    """用户点名“柱状图/折线图/饼图”时，优先于按类目数的启发式。"""
    columns = [
        {"name": "所属区域", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
    ]
    rows = [
        {"所属区域": "A区", "工单量": 420},
        {"所属区域": "B区", "工单量": 360},
        {"所属区域": "C区", "工单量": 139},
    ]

    assert infer_chart_meta(columns, rows, question="按所属区域统计，用柱状图展示").type == "bar"
    assert infer_chart_meta(columns, rows, question="按所属区域统计").type == "pie"
    assert (
        infer_chart_meta(
            columns,
            rows,
            question="先看饼图还是折线图？请用折线图",
        ).type
        == "line"
    )


def test_infer_chart_meta_uses_first_series_for_explicit_pie() -> None:
    """用户显式要饼图时，取第一个数值列，而不是降级成柱状图。"""
    columns = [
        {"name": "所属区域", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
        {"name": "总金额", "dtype": "float64"},
    ]
    rows = [{"所属区域": "A区", "工单量": 420, "总金额": 1000.0}]

    meta = infer_chart_meta(columns, rows, question="用饼图展示")

    assert meta.type == "pie"
    assert meta.y == ["工单量"]
    assert [series.name for series in meta.series] == ["工单量"]


def test_infer_chart_meta_falls_back_to_table_without_numeric_column() -> None:
    meta = infer_chart_meta(
        [{"name": "项目名称", "dtype": "object"}],
        [{"项目名称": "望京国际研发园"}],
    )

    assert meta.type == "table"
    assert meta.y == []


def test_export_writes_styled_excel_and_blocks_formula_injection(
    tmp_path: Path,
) -> None:
    exporter = ExcelReportExporter(_settings(tmp_path))
    exported = exporter.export(
        columns=[
            {"name": "所属区域", "dtype": "object"},
            {"name": "备注", "dtype": "object"},
            {"name": "工单量", "dtype": "int64"},
        ],
        rows=[
            {"所属区域": "朝阳区", "备注": "=1+1", "工单量": 132},
            {"所属区域": "海淀区", "备注": None, "工单量": 12},
        ],
    )

    assert exported is not None
    assert exported.path.is_file()
    assert exported.path.name.startswith("report_")
    assert exported.path.suffix == ".xlsx"

    workbook = load_workbook(exported.path, read_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()

    assert values[0] == ("所属区域", "备注", "工单量")
    assert values[1] == ("朝阳区", "'=1+1", 132)
    assert values[2][1] is None


def test_resolve_rejects_invalid_name_and_expired_file(tmp_path: Path) -> None:
    exporter = ExcelReportExporter(_settings(tmp_path))
    exported = exporter.export(
        columns=[{"name": "项目", "dtype": "object"}],
        rows=[{"项目": "A"}],
    )
    assert exported is not None

    assert exporter.resolve(exported.filename) == exported.path

    with pytest.raises(ReportNameInvalidError):
        exporter.resolve("../../etc/passwd")
    with pytest.raises(ReportNameInvalidError):
        exporter.resolve("report_not-a-token.xlsx")

    old = time.time() - 7200
    os.utime(exported.path, (old, old))
    with pytest.raises(ReportNotFoundError):
        exporter.resolve(exported.filename)
    assert not exported.path.exists()
