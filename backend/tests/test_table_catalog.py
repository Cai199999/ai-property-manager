import json
from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.services.mock_tables import generate_mock_tables
from app.services.table_catalog import TableCatalog


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        data_table_dir=tmp_path,
        table_directory_enabled=True,
        table_max_files=8,
        table_max_rows=10_000,
        table_max_columns=100,
    )


def test_catalog_deduplicates_identical_files(tmp_path: Path) -> None:
    """内容完全相同的副本只能保留一份，否则同构表 concat 会放大金额。"""
    generate_mock_tables(tmp_path)
    source = tmp_path / "工单表_A区.xlsx"
    (tmp_path / "upload_copy_work_orders.xlsx").write_bytes(source.read_bytes())

    snapshot = TableCatalog(_settings(tmp_path)).scan()

    assert {table.name for table in snapshot.tables} == {
        "工单表_A区",
        "工单表_B区",
        "工单表_C区",
    }
    assert all("upload_copy" not in table.file_name for table in snapshot.tables)
    assert any("内容相同" in warning for warning in snapshot.warnings)
    assert snapshot.homogeneous_groups[0]["tables"] == [
        "工单表_A区",
        "工单表_B区",
        "工单表_C区",
    ]


def test_catalog_detects_homogeneous_tables(tmp_path: Path) -> None:
    summary = generate_mock_tables(tmp_path)
    catalog = TableCatalog(_settings(tmp_path))

    snapshot = catalog.scan()

    names = {"工单表_A区", "工单表_B区", "工单表_C区"}
    assert snapshot.has_tables
    assert snapshot.primary_table == "工单表_A区"
    assert {table.name for table in snapshot.tables} == names
    assert snapshot.warnings == ()

    for table in snapshot.tables:
        assert table.role == "fact"
        assert table.schema_group == "work_orders"
        assert table.row_count == summary["tables"][table.name]
        # 原始 46 列 + 内存派生的费用金额。
        assert table.column_count == 47
        assert "费用金额" in table.frame.columns
        assert table.primary_key == ("单号",)
        assert table.time_column == "登记时间"
        assert table.foreign_keys == ()
        # 同构表之间不能生成 join 提示，否则模型可能按业务列做笛卡尔积。
        assert table.possible_join_keys == ()

    assert snapshot.relationships == ()
    assert len(snapshot.homogeneous_groups) == 1
    group = snapshot.homogeneous_groups[0]
    assert group["name"] == "work_orders"
    assert group["tables"] == ["工单表_A区", "工单表_B区", "工单表_C区"]
    assert group["concat_recommended"] is True
    assert len(group["columns"]) == 47


def test_catalog_uses_upload_index_for_display_name(tmp_path: Path) -> None:
    summary = generate_mock_tables(tmp_path)
    source = tmp_path / "工单表_A区.xlsx"
    target = tmp_path / "upload_abc123_work_orders.xlsx"
    # 上传件必须是不同内容，否则会命中“重复文件去重”而保留 manifest 原始表。
    book = load_workbook(source)
    book.worksheets[0].cell(row=3, column=1).value = "UPLOAD-COPY-001"
    book.save(target)
    (tmp_path / "upload_index.json").write_text(
        json.dumps({"upload_abc123_work_orders.xlsx": {"original_stem": "九月工单"}}),
        encoding="utf-8",
    )

    snapshot = TableCatalog(_settings(tmp_path)).scan()

    assert "九月工单" in snapshot.table_map()
    uploaded = snapshot.table_map()["九月工单"]
    assert uploaded.row_count == summary["tables"]["工单表_A区"]
    # 没有 manifest 声明的上传件不会被自动赋予外键。
    assert uploaded.foreign_keys == ()
    # 同构关系由列签名自动识别，上传件也应归入同一组。
    assert any("九月工单" in group["tables"] for group in snapshot.homogeneous_groups)


def test_catalog_skips_corrupt_file_without_failing(tmp_path: Path) -> None:
    generate_mock_tables(tmp_path)
    (tmp_path / "损坏表.xlsx").write_bytes(b"not an excel file")

    snapshot = TableCatalog(_settings(tmp_path)).scan()

    assert snapshot.has_tables
    assert any("损坏表.xlsx" in warning for warning in snapshot.warnings)
