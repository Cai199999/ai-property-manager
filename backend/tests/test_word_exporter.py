"""Word 报告导出与 AI 结论生成的单元测试。"""

import asyncio
import json
import zipfile
from pathlib import Path

import httpx
import pytest
from docx import Document
from pydantic import SecretStr

from app.config import Settings
from app.services.word_exporter import (
    LLMReportSummarizer,
    ReportSummary,
    WordReportExporter,
    build_fallback_summary,
    build_summary_messages,
    render_chart_png,
)


def _settings(tmp_path: Path, **overrides: object) -> Settings:
    values: dict[str, object] = {
        "report_export_enabled": True,
        "word_report_enabled": True,
        "report_export_dir": tmp_path / "exports",
        "report_export_ttl_seconds": 3600,
        "deepseek_api_key": SecretStr(""),
    }
    values.update(overrides)
    return Settings(_env_file=None, **values)


def _columns() -> list[dict[str, str]]:
    return [
        {"name": "所属区域", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
    ]


def _rows() -> list[dict[str, object]]:
    return [
        {"所属区域": "A区", "工单量": 420},
        {"所属区域": "B区", "工单量": 360},
        {"所属区域": "C区", "工单量": 139},
    ]


def _chart_meta() -> dict[str, object]:
    return {
        "type": "bar",
        "title": "工单量按所属区域",
        "x": "所属区域",
        "y": ["工单量"],
        "series": [{"name": "工单量", "key": "工单量", "unit": "单"}],
    }


def _summary() -> ReportSummary:
    return ReportSummary(text="本次查询共返回 3 行结果，A 区工单量最高。", engine="deepseek")


def _document_text(document: Document) -> str:
    pieces = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                pieces.append(cell.text)
    return "\n".join(pieces)


def test_export_builds_docx_with_all_sections(tmp_path: Path) -> None:
    exporter = WordReportExporter(_settings(tmp_path))

    exported = exporter.export(
        question="统计各区域工单量",
        columns=_columns(),
        rows=_rows(),
        chart_meta=_chart_meta(),
        summary=_summary(),
        request_id="word-test-001",
        time_range={"start_date": "2026-09-01", "end_date": "2026-09-30"},
        data_source="directory",
        table_names=("工单表_A区", "工单表_B区"),
        warnings=("示例告警",),
    )

    assert exported is not None
    assert exported.path.is_file()
    assert exported.path.name.startswith("report_")
    assert exported.path.suffix == ".docx"
    assert exported.path.stat().st_size > 10_000
    assert exported.path.stat().st_mode & 0o777 == 0o600

    document = Document(exported.path)
    text = _document_text(document)

    assert "物业工单分析报告" in text
    assert "统计各区域工单量" in text
    assert "本次查询共返回 3 行结果，A 区工单量最高。" in text
    assert "AI 生成" in text
    assert "2026-09-01 至 2026-09-30" in text
    assert "word-test-001" in text
    assert "示例告警" in text
    # 元信息表 + 数据明细表
    assert len(document.tables) == 2
    assert [cell.text for cell in document.tables[1].rows[0].cells] == ["所属区域", "工单量"]
    assert len(document.tables[1].rows) == 4
    # 折线/柱状图会插入一张 PNG
    assert len(document.inline_shapes) == 1

    with zipfile.ZipFile(exported.path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media, "Word 报告应包含图表图片"


def test_export_skips_when_disabled_or_empty(tmp_path: Path) -> None:
    disabled = WordReportExporter(_settings(tmp_path, word_report_enabled=False))
    assert (
        disabled.export(
            question="问题",
            columns=_columns(),
            rows=_rows(),
            chart_meta=_chart_meta(),
            summary=_summary(),
        )
        is None
    )

    enabled = WordReportExporter(_settings(tmp_path))
    assert (
        enabled.export(
            question="问题",
            columns=_columns(),
            rows=[],
            chart_meta=_chart_meta(),
            summary=_summary(),
        )
        is None
    )
    assert not (tmp_path / "exports").exists()


def test_export_truncates_rows_and_columns(tmp_path: Path) -> None:
    exporter = WordReportExporter(
        _settings(tmp_path, report_word_max_rows=2, report_word_max_columns=1)
    )
    rows = [{"项目": f"项目-{index}", "工单量": index} for index in range(1, 6)]
    columns = [
        {"name": "项目", "dtype": "object"},
        {"name": "工单量", "dtype": "int64"},
    ]

    exported = exporter.export(
        question="项目工单量",
        columns=columns,
        rows=rows,
        chart_meta={"type": "table"},
        summary=_summary(),
    )

    assert exported is not None
    document = Document(exported.path)
    data_table = document.tables[1]
    assert len(data_table.columns) == 1
    assert len(data_table.rows) == 3  # 表头 + 2 行
    assert "完整数据请下载 Excel" in _document_text(document)
    assert not document.inline_shapes


def test_export_without_chart_for_table_meta(tmp_path: Path) -> None:
    exporter = WordReportExporter(_settings(tmp_path))

    exported = exporter.export(
        question="按项目列出明细",
        columns=_columns(),
        rows=_rows(),
        chart_meta={"type": "table"},
        summary=ReportSummary(text="系统摘要", engine="fallback"),
    )

    assert exported is not None
    document = Document(exported.path)
    assert not document.inline_shapes
    assert "系统自动摘要" in _document_text(document)


@pytest.mark.parametrize("chart_type", ["bar", "line", "pie"])
def test_render_chart_png_writes_png(tmp_path: Path, chart_type: str) -> None:
    target = tmp_path / f"{chart_type}.png"
    meta = {**_chart_meta(), "type": chart_type}

    render_chart_png(
        columns=_columns(),
        rows=_rows(),
        chart_meta=meta,
        path=target,
        max_categories=12,
    )

    assert target.is_file()
    assert target.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
    assert target.stat().st_size > 5_000


def test_render_chart_rejects_mismatched_meta(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        render_chart_png(
            columns=_columns(),
            rows=_rows(),
            chart_meta={"type": "bar", "x": "不存在的列", "y": ["工单量"]},
            path=tmp_path / "x.png",
            max_categories=12,
        )


def test_build_fallback_summary_uses_result_metadata() -> None:
    text = build_fallback_summary(
        columns=_columns(),
        rows=_rows(),
        chart_meta=_chart_meta(),
    )

    assert "3 行" in text
    assert "所属区域" in text
    assert "工单量按所属区域" in text

    empty = build_fallback_summary(columns=_columns(), rows=[], chart_meta=None)
    assert "没有返回数据" in empty


def test_build_summary_messages_limits_context() -> None:
    columns = [{"name": f"列{index}", "dtype": "object"} for index in range(40)]
    rows = [{f"列{index}": "长文本" * 50 for index in range(40)} for _ in range(50)]

    messages = build_summary_messages(
        question="问题",
        columns=columns,
        rows=rows,
        chart_meta=_chart_meta(),
        time_range={"start_date": "2026-09-01", "end_date": "2026-09-30"},
        table_names=["工单表_A区"],
    )

    assert [message["role"] for message in messages] == ["system", "user"]
    payload = messages[1]["content"]
    context = json.loads(payload.split("<report_context>\n")[1].split("\n</report_context>")[0])
    assert len(context["columns"]) == 30
    assert len(context["rows_sample"]) == 20
    assert context["result_row_count"] == 50
    assert context["source_tables"] == ["工单表_A区"]
    # 单元格文本必须被截断，避免长文本撑爆上下文
    assert all(len(str(value)) <= 61 for value in context["rows_sample"][0].values())


def test_summarizer_returns_fallback_without_key(tmp_path: Path) -> None:
    summarizer = LLMReportSummarizer(_settings(tmp_path))

    summary = asyncio.run(
        summarizer.summarize(
            question="统计各区域工单量",
            columns=_columns(),
            rows=_rows(),
            chart_meta=_chart_meta(),
        )
    )
    asyncio.run(summarizer.aclose())

    assert summary.engine == "fallback"
    assert "3 行" in summary.text


def _mock_client(payload: object, status: int = 200) -> httpx.AsyncClient:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path.endswith("/chat/completions")
        assert request.headers["Authorization"] == "Bearer test-key"
        return httpx.Response(status, json=payload)

    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


def test_summarizer_returns_model_conclusion(tmp_path: Path) -> None:
    client = _mock_client(
        {
            "choices": [
                {"message": {"content": "A 区工单量 420 单，占全部结果的 46%，建议优先巡检。"}}
            ]
        }
    )
    summarizer = LLMReportSummarizer(
        _settings(tmp_path, deepseek_api_key=SecretStr("test-key")),
        client=client,
    )

    summary = asyncio.run(
        summarizer.summarize(
            question="统计各区域工单量",
            columns=_columns(),
            rows=_rows(),
            chart_meta=_chart_meta(),
        )
    )
    asyncio.run(client.aclose())

    assert summary.engine == "deepseek"
    assert summary.text.startswith("A 区工单量 420 单")


def test_summarizer_falls_back_on_upstream_error(tmp_path: Path) -> None:
    client = _mock_client({"error": {"message": "boom"}}, status=500)
    summarizer = LLMReportSummarizer(
        _settings(tmp_path, deepseek_api_key=SecretStr("test-key")),
        client=client,
    )

    summary = asyncio.run(
        summarizer.summarize(
            question="统计各区域工单量",
            columns=_columns(),
            rows=_rows(),
            chart_meta=_chart_meta(),
        )
    )
    asyncio.run(client.aclose())

    assert summary.engine == "fallback"
    assert "3 行" in summary.text
