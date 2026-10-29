import asyncio
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.adapters.base import FetchRequest
from app.adapters.rpa_fetcher import RPAFetcher
from app.domain.work_orders import REQUIRED_WORK_ORDER_COLUMNS, WorkOrderColumn
from app.schemas.assistant import TimeRange


def test_rpa_fetcher_downloads_excel_and_normalizes(tmp_path: Path) -> None:
    fetcher = RPAFetcher(tmp_path, simulate_delay_seconds=0)
    request = FetchRequest(
        time_range=TimeRange(start_date=date(2026, 9, 1), end_date=date(2026, 9, 5)),
        property_name="望京国际研发园",
    )

    result = asyncio.run(fetcher.fetch(request))

    assert result.source == "rpa"
    assert not result.frame.empty
    assert set(REQUIRED_WORK_ORDER_COLUMNS).issubset(result.frame.columns)
    assert result.frame[WorkOrderColumn.PROJECT.value].eq("望京国际研发园").all()
    assert {"安防告警", "工程维修"}.issubset(
        set(result.frame[WorkOrderColumn.REPAIR_TYPE.value].unique())
    )

    received = pd.to_datetime(result.frame[WorkOrderColumn.RECEIVED_AT.value])
    assert received.min().date() >= date(2026, 9, 1)
    assert received.max().date() <= date(2026, 9, 5)

    exported = list(tmp_path.glob("*.xlsx"))
    assert len(exported) == 1
    assert result.metadata["source_file"] == exported[0].name

    # 导出件必须保持真实报表的标题行、46 列表头和合并单元格版式。
    workbook = load_workbook(exported[0], data_only=True)
    worksheet = workbook["Sheet0"]
    assert worksheet["A1"].value == "工单汇总查询"
    assert "A1:AT1" in {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
    headers = [worksheet.cell(row=2, column=index).value for index in range(1, 47)]
    assert headers[:4] == ["单号", "登记时间", "项目名称", "工单类型"]
    assert headers[-1] == "未完成原因"
    assert worksheet.max_row > 3


def test_rpa_fetcher_passes_project_filter_to_project_column(tmp_path: Path) -> None:
    fetcher = RPAFetcher(tmp_path, simulate_delay_seconds=0)
    request = FetchRequest(
        time_range=TimeRange(start_date=date(2026, 9, 1), end_date=date(2026, 9, 2)),
        property_name="望京生命科学园",
    )

    result = asyncio.run(fetcher.fetch(request))

    assert result.frame[WorkOrderColumn.PROJECT.value].unique().tolist() == ["望京生命科学园"]
