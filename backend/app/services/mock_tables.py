"""生成与真实报表同构的多文件 Mock Excel。

真实场景下，用户会分多次导出不同区域/项目的工单，每个文件都是
「标题行 + 46 列表头 + 明细」的同一版式。本生成器据此产出 3 个结构
完全一致的文件，用于验证 ``pd.concat`` 纵向合并，而不是外键 JOIN。
"""

from __future__ import annotations

import json
import random
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.adapters.report_schema import (
    REPORT_COLUMN_WIDTHS,
    REPORT_COLUMNS,
    REPORT_SHEET_NAME,
    REPORT_TITLE,
)
from app.adapters.rpa_driver import SimulatedRpaExportDriver

MOCK_SEED = 20261029
SCHEMA_GROUP = "work_orders"

# 每个文件模拟一个区域，行数不同但表头完全一致。
REGION_SPECS: tuple[dict[str, Any], ...] = (
    {
        "name": "工单表_A区",
        "prefix": "A",
        "projects": ("望京国际研发园", "望京科技园", "望京数字创意园"),
        "rows": 420,
    },
    {
        "name": "工单表_B区",
        "prefix": "B",
        "projects": ("望京生命科学园", "望京未来信息园"),
        "rows": 360,
    },
    {
        "name": "工单表_C区",
        "prefix": "C",
        "projects": ("慧谷时空", "金色家园", "东湖街道办事处"),
        "rows": 300,
    },
)

LEGACY_FILES = ("工单表.xlsx", "项目表.xlsx", "人员表.xlsx")


def generate_mock_tables(
    output_dir: Path,
    *,
    seed: int = MOCK_SEED,
) -> dict[str, Any]:
    """清理旧三表并生成 3 个同构工单文件。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    for name in LEGACY_FILES:
        (output_dir / name).unlink(missing_ok=True)

    rng = random.Random(seed)
    driver = SimulatedRpaExportDriver()
    # 先构造一个完整的 9 月数据集，再按区域抽取不同行，保证字段格式、
    # 缺失值和状态分布与真实导出件一致。
    base = driver._build_frame(  # noqa: SLF001 - Mock 生成器复用同一套版式
        start_date=date(2026, 9, 1),
        end_date=date(2026, 9, 30),
        property_name=None,
    )

    tables: dict[str, int] = {}
    manifest_tables: list[dict[str, Any]] = []
    for spec in REGION_SPECS:
        frame = _sample_region_frame(base, spec, rng)
        name = str(spec["name"])
        target = output_dir / f"{name}.xlsx"
        write_table_workbook(
            frame,
            target,
            title=REPORT_TITLE,
            column_widths=REPORT_COLUMN_WIDTHS,
        )
        tables[name] = len(frame)
        manifest_tables.append(
            {
                "name": name,
                "file": target.name,
                "role": "fact",
                "schema_group": SCHEMA_GROUP,
                "description": (
                    f"{spec['prefix']}区工单明细；表头与原始报表完全一致，"
                    "多文件分析时必须按行纵向拼接"
                ),
                "primary_key": ["单号"],
                "time_column": "登记时间",
                "concat_tables": [item["name"] for item in REGION_SPECS],
            }
        )

    names = [item["name"] for item in REGION_SPECS]
    manifest = {
        "version": 2,
        "seed": seed,
        "schema_groups": [
            {
                "name": SCHEMA_GROUP,
                "tables": names,
                "columns": list(REPORT_COLUMNS),
                "concat_recommended": True,
                "description": "同构工单表，使用 pd.concat 纵向合并，禁止 pd.merge",
            }
        ],
        "tables": manifest_tables,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    # 旧上传索引指向已删除的 Mock 文件，重置为空。
    (output_dir / "upload_index.json").write_text("{}\n", encoding="utf-8")

    return {
        "output_dir": str(output_dir),
        "tables": tables,
        "total_work_orders": sum(tables.values()),
        "columns_per_table": len(REPORT_COLUMNS),
    }


def _sample_region_frame(
    base: pd.DataFrame,
    spec: dict[str, Any],
    rng: random.Random,
) -> pd.DataFrame:
    """从全量数据中按区域抽取行，并加区域前缀保证单号全局唯一。"""
    projects = tuple(spec["projects"])
    candidates = base.loc[base["项目名称"].isin(projects)].copy()
    if candidates.empty:
        raise RuntimeError(f"区域 {spec['name']} 没有可抽取的工单数据")

    target_rows = min(int(spec["rows"]), len(candidates))
    chosen: list[int] = []
    used: set[int] = set()

    # 每个项目至少保留少量工单，避免小项目在随机抽样中完全消失。
    for project in projects:
        project_indexes = candidates.index[candidates["项目名称"] == project].tolist()
        for index in project_indexes[:5]:
            chosen.append(index)
            used.add(index)

    remaining = [index for index in candidates.index.tolist() if index not in used]
    need = max(0, target_rows - len(chosen))
    if need:
        chosen.extend(rng.sample(remaining, k=min(need, len(remaining))))

    frame = candidates.loc[sorted(chosen[:target_rows])].copy()
    # 区域前缀让三个文件的单号不重复，便于对照 nunique 结果。
    frame["单号"] = frame["单号"].map(lambda value: f"{spec['prefix']}-{value}")
    frame = frame.loc[:, list(REPORT_COLUMNS)]
    return frame.sort_values("登记时间", ascending=False, kind="stable").reset_index(drop=True)


def write_table_workbook(
    frame: pd.DataFrame,
    target: Path,
    *,
    title: str,
    column_widths: dict[str, float] | None = None,
    sheet_name: str = REPORT_SHEET_NAME,
) -> None:
    """按真实导出件版式写表：标题行 + 表头行 + 明细。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_font = Font(name="Calibri", size=11)
    data_font = Font(name="Calibri", size=11)

    columns = [str(column) for column in frame.columns]
    last_column = get_column_letter(len(columns))
    worksheet.merge_cells(f"A1:{last_column}1")
    title_cell = worksheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30

    for column_index, column_name in enumerate(columns, start=1):
        cell = worksheet.cell(row=2, column=column_index, value=column_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        width = (column_widths or {}).get(column_name)
        if width is None:
            width = max(12.0, min(30.0, len(column_name) * 2.0 + 4.0))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.row_dimensions[2].height = 20

    for row_index, record in enumerate(frame.to_dict(orient="records"), start=3):
        for column_index, column_name in enumerate(columns, start=1):
            value = record.get(column_name)
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=_excel_value(value),
            )
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
        worksheet.row_dimensions[row_index].height = 20

    workbook.properties.creator = "Apache POI"
    workbook.save(target)


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            pass
    return value
