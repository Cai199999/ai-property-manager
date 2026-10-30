"""Word 可视化报告：AI 结论、数据表格与 matplotlib 静态图表。

报告文件与 Excel 共用 ``.runtime/exports`` 目录，由 Java 网关统一代理下载。
本模块只负责"总结文案 + 排版"，不执行生成代码、不接触数据库。
"""

from __future__ import annotations

import json
import logging
import math
import tempfile
import threading
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterator, Mapping, Sequence

import httpx
import matplotlib

# 服务端无显示环境，必须在 pyplot 导入前锁定 Agg 后端。
matplotlib.use("Agg")

import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib import font_manager  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from app.config import Settings  # noqa: E402
from app.services.llm_provider import (  # noqa: E402
    LLMProviderFactory,
    ResolvedProvider,
)
from app.services.report_export import (  # noqa: E402
    ExportFile,
    cleanup_report_directory,
)

logger = logging.getLogger(__name__)

# 与前端图表保持一致的系列色，避免报告和页面观感割裂。
_SERIES_COLORS = ("#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4")
_TITLE_COLOR = RGBColor(0x1B, 0x2A, 0x3A)
_MUTED_COLOR = RGBColor(0x6B, 0x72, 0x80)
_ACCENT_COLOR = RGBColor(0x2A, 0x78, 0xD6)
_HEADER_FILL = "DDEBF7"
_LABEL_FILL = "F2F5F9"
_ZEBRA_FILL = "FAFBFC"

_CJK_FONT_CANDIDATES = (
    "Noto Sans CJK SC",
    "Noto Sans CJK JP",
    "WenQuanYi Micro Hei",
    "Microsoft YaHei",
    "SimHei",
    "Droid Sans Fallback",
    "AR PL UMing CN",
)
_BODY_FONT = "Noto Sans CJK SC"

_SUMMARY_MAX_ROWS = 20
_SUMMARY_MAX_COLUMNS = 30
_SUMMARY_CELL_CHARS = 60

# matplotlib 不是线程安全的：导出任务跑在线程池里，绘图阶段加锁串行。
_PLOT_LOCK = threading.Lock()

_SUMMARY_SYSTEM_PROMPT = """你是资深物业运营数据分析师，负责为管理层撰写数据报告结论。
严格要求：
1. 只输出 150-250 字中文纯文本，可分 2-3 个自然段；禁止 Markdown、标题、编号列表和代码块。
2. 必须引用结果中的关键数字（结果行数、合计、占比、排名第一的类别等），数字只能来自用户提供的 report_context。
3. 先给总体结论，再指出 1-2 个值得关注的细节或可执行建议。
4. 信息不足时明确写“仅基于本次查询结果”，禁止编造字段、项目、金额或趋势。
5. 不要复述用户问题，不要输出“根据您的要求”等客套话。"""


@dataclass(frozen=True, slots=True)
class ReportSummary:
    """一段可直接写入 Word 的结论文案。"""

    text: str
    engine: str  # deepseek | doubao | qwen | fallback


class LLMReportSummarizer:
    """单独调用一次大模型生成报告结论；失败时退回规则摘要。"""

    def __init__(
        self,
        settings: Settings,
        *,
        client: httpx.AsyncClient | None = None,
        provider_factory: LLMProviderFactory | None = None,
    ) -> None:
        self._settings = settings
        self._providers = provider_factory or LLMProviderFactory(settings, client=client)
        self._owns_providers = provider_factory is None

    async def aclose(self) -> None:
        if self._owns_providers:
            await self._providers.aclose()

    async def summarize(
        self,
        *,
        question: str,
        columns: Sequence[Mapping[str, Any]],
        rows: Sequence[Mapping[str, Any]],
        chart_meta: Mapping[str, Any] | None,
        time_range: Mapping[str, Any] | None = None,
        table_names: Sequence[str] = (),
        provider: ResolvedProvider | None = None,
    ) -> ReportSummary:
        fallback = ReportSummary(
            text=build_fallback_summary(columns=columns, rows=rows, chart_meta=chart_meta),
            engine="fallback",
        )
        try:
            resolved = provider or self._providers.resolve()
        except Exception as exc:  # noqa: BLE001 - 总结属于可选增强，失败不能阻断报告
            logger.warning("Word 报告总结不可用: %s", type(exc).__name__)
            return fallback

        messages = build_summary_messages(
            question=question,
            columns=columns,
            rows=rows,
            chart_meta=chart_meta,
            time_range=time_range,
            table_names=table_names,
        )
        try:
            content = await self._providers.chat(
                resolved,
                messages,
                temperature=0.3,
                max_tokens=600,
                timeout=self._settings.report_summary_timeout_seconds,
            )
        except Exception as exc:  # noqa: BLE001 - 外部模型边界，失败即降级
            logger.warning(
                "Word 报告总结生成失败 provider=%s key_source=%s error=%s",
                resolved.provider_type,
                resolved.key_source,
                type(exc).__name__,
            )
            return fallback

        text = _clean_summary(content)
        if not text:
            logger.warning("Word 报告总结为空，使用规则摘要")
            return fallback
        return ReportSummary(text=text, engine=resolved.provider_type)


def build_summary_messages(
    *,
    question: str,
    columns: Sequence[Mapping[str, Any]],
    rows: Sequence[Mapping[str, Any]],
    chart_meta: Mapping[str, Any] | None,
    time_range: Mapping[str, Any] | None,
    table_names: Sequence[str],
) -> list[dict[str, str]]:
    """裁剪结果后拼装总结 Prompt，避免把全量明细发给模型。"""
    names = [str(column.get("name", "")) for column in columns][:_SUMMARY_MAX_COLUMNS]
    sample_rows = [
        {name: _summary_value(row.get(name)) for name in names} for row in rows[:_SUMMARY_MAX_ROWS]
    ]
    context = {
        "question": question,
        "result_row_count": len(rows),
        "columns": names,
        "chart": {
            "type": (chart_meta or {}).get("type", "table"),
            "title": (chart_meta or {}).get("title", ""),
            "x": (chart_meta or {}).get("x"),
            "y": (chart_meta or {}).get("y", []),
        },
        "time_range": dict(time_range or {}),
        "source_tables": list(table_names),
        "rows_sample": sample_rows,
    }
    return [
        {"role": "system", "content": _SUMMARY_SYSTEM_PROMPT},
        {
            "role": "user",
            "content": (
                "下面是本次查询的上下文，请据此撰写报告结论。"
                "其中 rows_sample 只是结果的前若干行，不等同于全部明细。\n"
                "<report_context>\n"
                f"{json.dumps(context, ensure_ascii=False, default=str)}\n"
                "</report_context>"
            ),
        },
    ]


def build_fallback_summary(
    *,
    columns: Sequence[Mapping[str, Any]],
    rows: Sequence[Mapping[str, Any]],
    chart_meta: Mapping[str, Any] | None,
) -> str:
    """模型不可用时的确定性摘要，不臆造任何数字。"""
    names = [str(column.get("name", "")) for column in columns][:5]
    if not rows:
        return "本次查询没有返回数据，建议放宽筛选条件后重试。完整结果请查看随附的 Excel 文件。"

    parts = [f"本次查询共返回 {len(rows):,} 行结果。"]
    if names:
        parts.append(f"结果包含{'、'.join(names)}等字段。")
    chart_type = (chart_meta or {}).get("type", "table")
    if chart_type in {"bar", "line", "pie"}:
        title = (chart_meta or {}).get("title") or "查询结果"
        parts.append(f"报告已根据结果生成“{title}”图表。")
    parts.append("完整明细请查看报告中的数据表格及随附的 Excel 文件。")
    return "".join(parts)


class WordReportExporter:
    """把查询结果排版成 docx，并保证文件落在受控目录内。"""

    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._directory = Path(settings.report_export_dir)

    @property
    def enabled(self) -> bool:
        return self._settings.report_export_enabled and self._settings.word_report_enabled

    def export(
        self,
        *,
        question: str,
        columns: list[dict[str, str]],
        rows: list[dict[str, Any]],
        chart_meta: Mapping[str, Any] | None,
        summary: ReportSummary,
        request_id: str = "",
        time_range: Mapping[str, Any] | None = None,
        data_source: str = "",
        table_names: Sequence[str] = (),
        warnings: Sequence[str] = (),
    ) -> ExportFile | None:
        """生成 Word 报告；空结果、非法列或超限时返回 None。"""
        if not self.enabled or not rows:
            return None

        names = [str(column.get("name", "")) for column in columns]
        if not names or len(set(names)) != len(names):
            logger.warning("Word 报告列定义不合法，跳过生成")
            return None

        self._directory.mkdir(parents=True, exist_ok=True)
        filename = f"report_{uuid.uuid4().hex}.docx"
        target = self._directory / filename
        temporary = self._directory / f".{filename}.part"

        try:
            with tempfile.TemporaryDirectory(prefix="word-report-") as tmp:
                chart_path = None
                if _can_render_chart(chart_meta, rows):
                    chart_path = Path(tmp) / "chart.png"
                    render_chart_png(
                        columns=columns,
                        rows=rows,
                        chart_meta=chart_meta,
                        path=chart_path,
                        max_categories=self._settings.report_chart_max_categories,
                    )
                document = self._build_document(
                    question=question,
                    columns=columns,
                    rows=rows,
                    chart_meta=chart_meta,
                    chart_path=chart_path,
                    summary=summary,
                    request_id=request_id,
                    time_range=time_range,
                    data_source=data_source,
                    table_names=table_names,
                    warnings=warnings,
                )
                document.save(temporary)

            size = temporary.stat().st_size
            if size > self._settings.report_export_max_bytes:
                temporary.unlink(missing_ok=True)
                logger.warning("Word 报告超过大小限制 size=%s", size)
                return None
            temporary.replace(target)
            target.chmod(0o600)
        except Exception:  # noqa: BLE001 - 导出失败不应影响查询结果返回
            temporary.unlink(missing_ok=True)
            target.unlink(missing_ok=True)
            logger.exception("Word 报告生成失败")
            return None

        cleanup_report_directory(
            self._directory,
            ttl_seconds=self._settings.report_export_ttl_seconds,
            max_files=self._settings.report_export_max_files,
        )
        return ExportFile(filename=filename, path=target, size_bytes=size)

    def _build_document(
        self,
        *,
        question: str,
        columns: list[dict[str, str]],
        rows: list[dict[str, Any]],
        chart_meta: Mapping[str, Any] | None,
        chart_path: Path | None,
        summary: ReportSummary,
        request_id: str,
        time_range: Mapping[str, Any] | None,
        data_source: str,
        table_names: Sequence[str],
        warnings: Sequence[str],
    ) -> Document:
        document = Document()
        _configure_document(document)
        generated_at = datetime.now().astimezone()

        _add_title(document, question=question)
        _add_meta_table(
            document,
            rows=[
                ("数据源", _source_label(data_source)),
                ("时间范围", _format_time_range(time_range)),
                ("结果规模", f"{len(rows):,} 行 × {len(columns)} 列"),
                ("涉及数据表", "、".join(table_names) or "—"),
                ("生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
                ("请求编号", request_id or "—"),
            ],
        )

        _add_heading(document, "一、用户提问")
        _add_body_paragraph(document, question or "—")

        _add_heading(document, "二、分析结论")
        # engine 为实际 Provider（deepseek/doubao/qwen）或 fallback；
        # 只要不是本地兜底摘要，都应标记为 AI 生成。
        label = "系统自动摘要" if summary.engine == "fallback" else "AI 生成"
        note = document.add_paragraph()
        note_run = note.add_run(f"以下结论由{label}，仅基于本次查询结果。")
        _style_run(note_run, size=9, color=_MUTED_COLOR)
        for paragraph_text in summary.text.splitlines():
            text = paragraph_text.strip()
            if text:
                _add_body_paragraph(document, text, line_spacing=1.5)

        _add_heading(document, "三、数据明细")
        truncated = _add_data_table(
            document,
            columns=columns,
            rows=rows,
            max_rows=self._settings.report_word_max_rows,
            max_columns=self._settings.report_word_max_columns,
        )
        if truncated:
            _add_note(
                document,
                f"表格仅展示前 {self._settings.report_word_max_rows} 行、"
                f"前 {self._settings.report_word_max_columns} 列；完整数据请下载 Excel。",
            )

        if chart_path is not None:
            _add_heading(document, "四、可视化图表")
            picture = document.add_picture(str(chart_path), width=Cm(16.0))
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title = (chart_meta or {}).get("title") or "查询结果图表"
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_run(caption.add_run(str(title)), size=9, color=_MUTED_COLOR)

        if warnings:
            _add_heading(document, "附：数据说明")
            for warning in warnings:
                _add_note(document, f"· {warning}")

        return document


def render_chart_png(
    *,
    columns: list[dict[str, str]],
    rows: list[dict[str, Any]],
    chart_meta: Mapping[str, Any] | None,
    path: Path,
    max_categories: int,
) -> None:
    """按 chart_meta 绘制静态图并保存 PNG（调用方负责临时目录清理）。"""
    _configure_matplotlib()
    meta = chart_meta or {}
    chart_type = str(meta.get("type") or "table")
    x_key = meta.get("x")
    y_keys = [str(key) for key in (meta.get("y") or [])]
    title = str(meta.get("title") or "查询结果")

    table = pd.DataFrame(
        rows,
        columns=[str(column.get("name", "")) for column in columns],
    )
    if x_key not in table.columns or not y_keys:
        raise ValueError("chart_meta 与结果列不匹配")

    labels = [_short_label(value) for value in table[x_key].tolist()][:max_categories]
    values = {
        key: pd.to_numeric(table[key], errors="coerce").fillna(0.0).tolist()[:max_categories]
        for key in y_keys
        if key in table.columns
    }

    # matplotlib 的字体缓存、rcParams 和 Glyph 缓存都是全局状态，并发导出必须串行。
    with _PLOT_LOCK:
        figure, axis = plt.subplots(figsize=(8.6, 4.4), dpi=160)
        try:
            if chart_type == "pie":
                _draw_pie(axis, labels=labels, values=values.get(y_keys[0], []))
            elif chart_type == "line":
                _draw_line(axis, labels=labels, values=values)
            else:
                _draw_bar(axis, labels=labels, values=values)
            axis.set_title(title, fontsize=13, fontweight="bold", color="#1b2a3a", pad=14)
            figure.tight_layout()
            # bbox_inches="tight" 保证右侧饼图图例不会被裁掉。
            figure.savefig(
                path,
                format="png",
                dpi=160,
                facecolor="#fcfcfb",
                bbox_inches="tight",
                pad_inches=0.18,
            )
        finally:
            plt.close(figure)


def _draw_bar(axis: Any, *, labels: list[str], values: dict[str, list[float]]) -> None:
    x_positions = range(len(labels))
    series = list(values.items())
    width = 0.8 / max(len(series), 1)
    for index, (name, data) in enumerate(series):
        offset = (index - (len(series) - 1) / 2) * width
        bars = axis.bar(
            [position + offset for position in x_positions],
            data,
            width=width * 0.92,
            label=name,
            color=_SERIES_COLORS[index % len(_SERIES_COLORS)],
            edgecolor="#fcfcfb",
            linewidth=1.2,
        )
        if len(series) == 1 and len(data) <= 12:
            axis.bar_label(bars, fmt="%g", padding=3, fontsize=8, color="#4b5563")
    axis.set_xticks(list(x_positions))
    axis.set_xticklabels(labels, rotation=30 if len(labels) > 5 else 0, ha="right")
    if len(series) > 1:
        axis.legend(frameon=False, fontsize=9)
    _style_axes(axis)


def _draw_line(axis: Any, *, labels: list[str], values: dict[str, list[float]]) -> None:
    positions = range(len(labels))
    for index, (name, data) in enumerate(values.items()):
        axis.plot(
            positions,
            data,
            marker="o",
            markersize=5,
            linewidth=2,
            label=name,
            color=_SERIES_COLORS[index % len(_SERIES_COLORS)],
        )
    axis.set_xticks(list(positions))
    axis.set_xticklabels(labels, rotation=30 if len(labels) > 8 else 0, ha="right")
    if len(values) > 1:
        axis.legend(frameon=False, fontsize=9)
    _style_axes(axis)


def _draw_pie(axis: Any, *, labels: list[str], values: Sequence[float]) -> None:
    total = sum(float(value) for value in values if value > 0)
    if total <= 0:
        axis.text(0.5, 0.5, "无数值可绘制", ha="center", va="center", color="#6b7280")
        axis.axis("off")
        return
    colors = [_SERIES_COLORS[index % len(_SERIES_COLORS)] for index in range(len(labels))]
    wedges, _, autotexts = axis.pie(
        values,
        colors=colors,
        startangle=90,
        counterclock=False,
        autopct=lambda percent: f"{percent:.0f}%" if percent >= 3 else "",
        pctdistance=0.78,
        wedgeprops={"width": 0.42, "edgecolor": "#fcfcfb", "linewidth": 2},
        textprops={"fontsize": 9, "color": "#374151"},
    )
    for text in autotexts:
        text.set_fontsize(8)
    axis.legend(
        wedges,
        labels,
        loc="center left",
        bbox_to_anchor=(1.0, 0.5),
        frameon=False,
        fontsize=9,
    )
    axis.set_aspect("equal")


def _configure_matplotlib() -> None:
    """锁定中文字体，避免 PNG 出现方框。"""
    selected = _resolve_cjk_font()
    if selected:
        plt.rcParams["font.sans-serif"] = [selected, "DejaVu Sans"]
    else:
        plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


@lru_cache(maxsize=1)
def _resolve_cjk_font() -> str | None:
    """优先使用已登记的中文字体；TTC 未登记时手动注册后再查一次。"""
    selected = _pick_installed_font()
    if selected:
        return selected

    for path in _iter_cjk_font_files():
        try:
            font_manager.fontManager.addfont(str(path))
        except (RuntimeError, OSError, ValueError):
            logger.warning("注册中文字体失败 path=%s", path)
            continue
        selected = _pick_installed_font()
        if selected:
            logger.info("已注册中文字体 name=%s path=%s", selected, path)
            return selected

    logger.warning("未找到可用的中文字体，图表文字可能显示为方框")
    return None


def _pick_installed_font() -> str | None:
    names = {font.name for font in font_manager.fontManager.ttflist}
    return next((name for name in _CJK_FONT_CANDIDATES if name in names), None)


def _iter_cjk_font_files() -> Iterator[Path]:
    """扫描常见字体目录，兼容 matplotlib 未解析 TTC 子字体的环境。"""
    patterns = ("*CJK*.ttc", "*CJK*.otf", "*SC*.otf", "*Hei*.ttf", "*Song*.ttf", "wqy*")
    for directory in (Path("/usr/share/fonts"), Path.home() / ".local/share/fonts"):
        if not directory.is_dir():
            continue
        for pattern in patterns:
            yield from sorted(directory.rglob(pattern))


def _style_axes(axis: Any) -> None:
    axis.grid(axis="y", color="#e5e7eb", linewidth=0.8)
    axis.set_axisbelow(True)
    for side in ("top", "right"):
        axis.spines[side].set_visible(False)
    axis.spines["left"].set_color("#d1d5db")
    axis.spines["bottom"].set_color("#d1d5db")
    axis.tick_params(colors="#6b7280", labelsize=9)


def _can_render_chart(chart_meta: Mapping[str, Any] | None, rows: Sequence[Any]) -> bool:
    meta = chart_meta or {}
    chart_type = meta.get("type")
    if chart_type not in {"bar", "line", "pie"}:
        return False
    return bool(meta.get("x")) and bool(meta.get("y")) and bool(rows)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)

    properties = document.core_properties
    properties.author = "AI 物管助手"
    properties.title = "物业工单分析报告"
    properties.subject = "自然语言查询结果可视化报告"


def _add_title(document: Document, *, question: str) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("物业工单分析报告")
    _style_run(run, size=24, bold=True, color=_TITLE_COLOR)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    _style_run(subtitle.add_run("AI Property Operations Report"), size=9, color=_MUTED_COLOR)

    if question:
        summary_line = document.add_paragraph()
        summary_line.paragraph_format.space_after = Pt(14)
        _style_run(summary_line.add_run(f"“{question}”"), size=11, color=_ACCENT_COLOR)

    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8")
    border.set(qn("w:space"), "6")
    border.set(qn("w:color"), "2A78D6")
    bottom.append(border)
    subtitle._p.get_or_add_pPr().append(bottom)


def _add_meta_table(document: Document, *, rows: Sequence[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table, color="DDE3EA")
    for index, (label, value) in enumerate(rows):
        left, right = table.rows[index].cells
        left.width = Cm(3.2)
        right.width = Cm(12.6)
        _fill_cell(left, _LABEL_FILL)
        _write_cell(left, label, bold=True, color=_TITLE_COLOR, size=9.5)
        _write_cell(right, value, color=RGBColor(0x37, 0x41, 0x51), size=9.5)
    document.add_paragraph()


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _style_run(run, size=14, bold=True, color=_TITLE_COLOR)


def _add_body_paragraph(
    document: Document,
    text: str,
    *,
    line_spacing: float = 1.3,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _style_run(paragraph.add_run(text), size=10.5, color=RGBColor(0x37, 0x41, 0x51))


def _add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    _style_run(paragraph.add_run(text), size=8.5, color=_MUTED_COLOR)


def _add_data_table(
    document: Document,
    *,
    columns: list[dict[str, str]],
    rows: list[dict[str, Any]],
    max_rows: int,
    max_columns: int,
) -> bool:
    """写入结果表；返回是否发生行列截断。"""
    names = [str(column.get("name", "")) for column in columns]
    selected = names[:max_columns]
    visible_rows = rows[:max_rows]

    table = document.add_table(rows=len(visible_rows) + 1, cols=len(selected))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table, color="E1E6EC")

    header = table.rows[0]
    _repeat_header(header)
    for index, name in enumerate(selected):
        cell = header.cells[index]
        _fill_cell(cell, _HEADER_FILL)
        _write_cell(cell, name, bold=True, color=_TITLE_COLOR, size=9)

    for row_index, row in enumerate(visible_rows, start=1):
        table_row = table.rows[row_index]
        if row_index % 2 == 0:
            for cell in table_row.cells:
                _fill_cell(cell, _ZEBRA_FILL)
        for column_index, name in enumerate(selected):
            _write_cell(
                table_row.cells[column_index],
                _display_value(row.get(name)),
                size=8.5,
                color=RGBColor(0x37, 0x41, 0x51),
            )

    document.add_paragraph()
    return len(rows) > max_rows or len(names) > max_columns


def _set_table_borders(table: Any, *, color: str) -> None:
    table_properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    table_properties.append(borders)


def _fill_cell(cell: Any, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _write_cell(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor | None = None,
    size: float = 9,
) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    _style_run(run, size=size, bold=bold, color=color)


def _repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _style_run(
    run: Any,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = _BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _BODY_FONT)


def _source_label(source: str) -> str:
    return {
        "directory": "本地数据目录（多个 Excel 文件）",
        "rpa": "RPA 导出数据",
        "api": "业务系统 API",
    }.get(source, source or "—")


def _format_time_range(time_range: Mapping[str, Any] | None) -> str:
    if not time_range:
        return "—"
    start = _display_value(time_range.get("start_date"))
    end = _display_value(time_range.get("end_date"))
    if start == "—" and end == "—":
        return "—"
    return f"{start} 至 {end}"


def _display_value(value: Any) -> str:
    """把 JSON 标量转成适合 Word 单元格的短文本。"""
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        if not math.isfinite(value):
            return "—"
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if hasattr(value, "item"):
        try:
            return _display_value(value.item())
        except (ValueError, AttributeError):
            pass
    text = str(value)
    return text[:80] + ("…" if len(text) > 80 else "")


def _summary_value(value: Any) -> Any:
    """总结 Prompt 只需要短标量，避免超长文本挤占 token。"""
    if value is None:
        return None
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            pass
    text = str(value)
    return text[:_SUMMARY_CELL_CHARS] + ("…" if len(text) > _SUMMARY_CELL_CHARS else "")


def _clean_summary(content: Any) -> str:
    if not isinstance(content, str):
        return ""
    text = content.strip()
    if text.startswith("```") and text.endswith("```"):
        text = text.strip("`").strip()
        if text.lower().startswith("text"):
            text = text[4:].strip()
    return text[:1_200]


def _short_label(value: Any) -> str:
    text = _display_value(value)
    return text[:14] + ("…" if len(text) > 14 else "")
