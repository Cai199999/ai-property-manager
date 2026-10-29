"""查询结果 Excel 导出、过期清理与安全下载解析。

沙箱只返回 JSON 安全的结果行，导出进程只负责把这些行写成 xlsx。文件使用
随机 token 命名，并限制在固定目录内，避免把服务器路径暴露给模型或调用方。
"""

from __future__ import annotations

import logging
import math
import re
import threading
import time
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import Settings
from app.core.exceptions import ReportNameInvalidError, ReportNotFoundError

logger = logging.getLogger(__name__)

REPORT_DOWNLOAD_PATH = "/api/v1/assistant/download"
REPORT_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WORD_REPORT_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# 下载路由按扩展名选择 MIME，Excel 与 Word 共用同一个导出目录和下载前缀。
REPORT_MEDIA_TYPES = {
    ".xlsx": REPORT_MEDIA_TYPE,
    ".docx": WORD_REPORT_MEDIA_TYPE,
}

_EXPORT_NAME_RE = re.compile(r"^report_[0-9a-f]{32}\.(?:xlsx|docx)$")
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


@dataclass(frozen=True, slots=True)
class ExportFile:
    """一次成功导出的文件元数据。"""

    filename: str
    path: Path
    size_bytes: int


def media_type_for(filename: str) -> str:
    """按扩展名返回下载 MIME；调用前应先通过文件名白名单校验。"""
    return REPORT_MEDIA_TYPES.get(Path(filename).suffix.lower(), REPORT_MEDIA_TYPE)


def cleanup_report_directory(
    directory: Path,
    *,
    ttl_seconds: int,
    max_files: int,
) -> None:
    """删除过期文件，并在数量超限时按修改时间淘汰最旧文件。"""
    try:
        files = sorted(
            (
                path
                for path in directory.glob("report_*")
                if path.suffix.lower() in REPORT_MEDIA_TYPES and path.is_file()
            ),
            key=lambda item: item.stat().st_mtime,
        )
    except OSError:
        return

    now = time.time()
    alive: list[Path] = []
    for path in files:
        try:
            if now - path.stat().st_mtime > ttl_seconds:
                path.unlink(missing_ok=True)
            else:
                alive.append(path)
        except OSError:
            continue

    excess = len(alive) - max_files
    for path in alive[: max(excess, 0)]:
        path.unlink(missing_ok=True)


class ExcelReportExporter:
    """把查询结果写成受目录、大小和 TTL 约束的 xlsx 文件。"""

    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._directory = Path(settings.report_export_dir)
        self._lock = threading.Lock()

    @property
    def enabled(self) -> bool:
        return self._settings.report_export_enabled

    def export(
        self,
        *,
        columns: list[dict[str, str]],
        rows: list[dict[str, Any]],
    ) -> ExportFile | None:
        """导出结果表；空结果或超过大小限制时返回 None。"""
        if not self.enabled or not rows:
            return None

        names = [str(column.get("name", "")) for column in columns]
        if not names or len(set(names)) != len(names):
            logger.warning("导出结果列定义不合法，跳过 Excel 导出")
            return None

        self._directory.mkdir(parents=True, exist_ok=True)
        filename = f"report_{uuid.uuid4().hex}.xlsx"
        target = self._directory / filename
        temporary = self._directory / f".{filename}.part"

        try:
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("查询结果")
            sheet.freeze_panes = "A2"
            for index, name in enumerate(names, start=1):
                width = min(max(len(name) * 2 + 4, 12), 40)
                sheet.column_dimensions[get_column_letter(index)].width = width

            header_cells = []
            for name in names:
                cell = WriteOnlyCell(sheet, value=_excel_value(name))
                cell.font = Font(bold=True)
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(horizontal="left", vertical="center")
                header_cells.append(cell)
            sheet.append(header_cells)

            for row in rows:
                sheet.append([_excel_value(row.get(name)) for name in names])

            workbook.save(temporary)
            size = temporary.stat().st_size
            if size > self._settings.report_export_max_bytes:
                temporary.unlink(missing_ok=True)
                logger.warning("导出文件超过大小限制 size=%s", size)
                return None
            temporary.replace(target)
            target.chmod(0o600)
        except Exception:  # noqa: BLE001 - 导出失败不应影响查询结果返回
            temporary.unlink(missing_ok=True)
            target.unlink(missing_ok=True)
            logger.exception("Excel 导出失败")
            return None

        with self._lock:
            self._cleanup_locked()
        return ExportFile(filename=filename, path=target, size_bytes=size)

    def resolve(self, filename: str) -> Path:
        """校验下载文件名并返回目录内的真实路径。"""
        if not self.enabled or not filename or not _EXPORT_NAME_RE.fullmatch(filename):
            raise ReportNameInvalidError()

        directory = self._directory.resolve()
        path = (directory / filename).resolve()
        if path.parent != directory:
            raise ReportNameInvalidError()

        try:
            modified_at = path.stat().st_mtime
        except FileNotFoundError as exc:
            raise ReportNotFoundError() from exc

        if time.time() - modified_at > self._settings.report_export_ttl_seconds:
            path.unlink(missing_ok=True)
            raise ReportNotFoundError()
        return path

    def _cleanup_locked(self) -> None:
        """Excel 与 Word 共用目录，清理时一并处理两种后缀。"""
        cleanup_report_directory(
            self._directory,
            ttl_seconds=self._settings.report_export_ttl_seconds,
            max_files=self._settings.report_export_max_files,
        )


def _excel_value(value: Any) -> Any:
    """转换单元格值，并阻断 Excel 公式注入。"""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (datetime, date, bool, int, float)):
        return value

    text = str(value)
    if text.startswith(_DANGEROUS_PREFIXES):
        return "'" + text
    return text
